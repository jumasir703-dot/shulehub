"""
import_files.py — turn an uploaded Excel or PDF file into plain rows of text.

The importers in this app were written around CSV. Rather than touch each of
them, `analytics._read_csv_upload` calls `rows_from_upload()` here for any
file that is not CSV, gets back a list of rows, and re-serialises them as CSV
text so the existing header detection and validation run unchanged.

    .xlsx / .xlsm   read with openpyxl (first sheet that has data)
    .pdf            tables read with pdfplumber (text layer only; scanned
                    PDFs have no text to read and are rejected with a hint)

Both libraries are imported lazily so the rest of the app still runs if one
is not installed.
"""

import io

MAX_BINARY_UPLOAD_BYTES = 10 * 1024 * 1024  # Excel/PDF files are bigger than CSV
MAX_PDF_PAGES = 60


def sniff_kind(data, filename=""):
    """Return 'xlsx', 'xls', 'pdf' or None (treat as CSV/text) from the bytes."""
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        return "xlsx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    return None


def _clean(value):
    """Cell -> string. Whole-number floats lose the '.0' (1001.0 -> 1001)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().replace("\r", " ").replace("\n", " ")


def _rows_from_xlsx(data):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None, "Excel import isn't installed on this server (missing openpyxl). Save the sheet as CSV instead."
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return None, ("That file couldn't be read as an Excel workbook. If it is password-protected "
                      "or damaged, re-save it as .xlsx or CSV and try again.")
    for ws in wb.worksheets:
        rows = [[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if rows:
            return rows, None
    return None, "The workbook has no data in it."


def _rows_from_pdf(data):
    try:
        import pdfplumber
    except ImportError:
        return None, "PDF import isn't installed on this server (missing pdfplumber). Use Excel or CSV instead."
    rows = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                return None, f"That PDF has {len(pdf.pages)} pages (limit {MAX_PDF_PAGES}). Split it and upload in parts."
            for page in pdf.pages:
                tables = page.extract_tables()
                if not tables:
                    # Tables drawn without ruling lines: fall back to aligned text.
                    tables = page.extract_tables({"vertical_strategy": "text", "horizontal_strategy": "text"})
                for table in tables:
                    for row in table:
                        cells = [_clean(c) for c in row]
                        if any(cells):
                            rows.append(cells)
    except Exception:
        return None, ("That PDF couldn't be read. It may be damaged or password-protected. "
                      "Open it, export the table to Excel, and upload that instead.")
    if not rows:
        return None, ("No table with text was found in that PDF. If it is a scan or a photo, it has no text "
                      "to read — upload the original Excel file or a CSV instead.")
    return _drop_repeated_headers(rows), None


def _drop_repeated_headers(rows):
    """A table that spans several pages repeats its heading row on each page."""
    header = rows[0]
    return [header] + [r for r in rows[1:] if r != header]


def rows_from_upload(data, kind):
    """-> (rows, error). `kind` comes from sniff_kind()."""
    if kind == "xlsx":
        return _rows_from_xlsx(data)
    if kind == "pdf":
        return _rows_from_pdf(data)
    if kind == "xls":
        return None, ("That is an old Excel (.xls) file. Open it in Excel, use File → Save As → "
                      "“Excel Workbook (.xlsx)”, and upload that.")
    return None, "Unsupported file type."
